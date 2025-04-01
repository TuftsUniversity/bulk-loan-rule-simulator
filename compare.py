import pandas as pd
from tkinter import Tk, filedialog
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
import os

# Highlight styles
YELLOW = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")  # changed cells
RED = PatternFill(start_color="FF9999", end_color="FF9999", fill_type="solid")      # missing rows

def select_excel_file(prompt):
    root = Tk()
    root.withdraw()
    return filedialog.askopenfilename(title=prompt, filetypes=[("Excel files", "*.xlsx")])

def compare_and_highlight(file1, file2):
    df1 = pd.read_excel(file1, dtype=str).fillna("")
    df2 = pd.read_excel(file2, dtype=str).fillna("")

    # Remove exact duplicates from both
    df1.drop_duplicates(inplace=True)
    df2.drop_duplicates(inplace=True)

    # Multi-level sort: User ID then Barcode
    def get_sort_cols(df):
        # Match columns case-insensitively
        lower_cols = [col.lower() for col in df.columns]
        user_id_col = next((c for c in df.columns if c.lower() == "user id"), df.columns[0])
        barcode_col = next((c for c in df.columns if c.lower() == "barcode"), df.columns[1])
        return [user_id_col, barcode_col]

    sort_cols = get_sort_cols(df1)  # same columns assumed in both
    df1.sort_values(by=sort_cols, inplace=True)
    df2.sort_values(by=sort_cols, inplace=True)
    df1.reset_index(drop=True, inplace=True)
    df2.reset_index(drop=True, inplace=True)

    # Pad to equal row counts
    max_rows = max(len(df1), len(df2))
    while len(df1) < max_rows:
        df1.loc[len(df1)] = [""] * len(df1.columns)
    while len(df2) < max_rows:
        df2.loc[len(df2)] = [""] * len(df2.columns)

    # # Validate column match
    # if list(df1.columns) != list(df2.columns):
    #     print("❌ Column mismatch between files.")
    #     return

    # Compare and write highlighted Excel output
    def write_with_highlights(df_base, df_compare, original_file, label):
        base_filename = os.path.basename(original_file)
        output_path = f"highlighted_{label}_{base_filename}"
        df_base.to_excel(output_path, index=False)

        wb = load_workbook(output_path)
        ws = wb.active

        for r in range(len(df_base)):
            base_row = df_base.iloc[r]
            compare_row = df_compare.iloc[r]

            if all(val == "" for val in compare_row):
                for c in range(len(df_base.columns)):
                    ws.cell(row=r+2, column=c+1).fill = RED
            else:
                for c, col in enumerate(df_base.columns):
                    if base_row[col] != compare_row[col]:
                        ws.cell(row=r+2, column=c+1).fill = YELLOW

        wb.save(output_path)
        print(f"✅ Saved highlighted file: {output_path}")

    write_with_highlights(df1, df2, file1, "file1")
    write_with_highlights(df2, df1, file2, "file2")


if __name__ == "__main__":
    file1 = select_excel_file("Select the FIRST Excel file")
    file2 = select_excel_file("Select the SECOND Excel file")

    if file1 and file2:
        compare_and_highlight(file1, file2)
    else:
        print("❌ File selection cancelled.")
